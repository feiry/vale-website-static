#!/usr/bin/env python3
"""Generate the COMMs Website Update Workflow rollout deck (.pptx).
Audience: COMMs team + IT/GDI. Purpose: roll out the request→publish workflow.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ── Vale brand palette ────────────────────────────────────────────────
TEAL   = RGBColor(0x00, 0x7A, 0x5E)   # vale green/teal
TEAL_D = RGBColor(0x00, 0x53, 0x40)
YELLOW = RGBColor(0xF2, 0xA9, 0x00)   # vale yellow
DARK   = RGBColor(0x2B, 0x2B, 0x2B)
GREY   = RGBColor(0x5A, 0x5A, 0x5A)
LIGHT  = RGBColor(0xF4, 0xF6, 0xF5)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.333)   # 16:9
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, fill, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    """runs: list of paragraphs; each paragraph is list of (text,size,color,bold) tuples."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(4)
        for (t, sz, col, bold) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = col; r.font.bold = bold
            r.font.name = "Calibri"
    return tb


def bullets(s, x, y, w, h, items, size=16, color=DARK, gap=8):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.05
        lvl = 0
        if isinstance(it, tuple):
            it, lvl = it
        r = p.add_run()
        bullet = "•  " if lvl == 0 else "–  "
        r.text = bullet + it
        r.font.size = Pt(size if lvl == 0 else size - 2)
        r.font.color.rgb = color if lvl == 0 else GREY
        r.font.name = "Calibri"
        p.level = lvl
    return tb


def header(s, title, kicker="WEBSITE UPDATE WORKFLOW"):
    rect(s, 0, 0, SW, Inches(1.15), TEAL)
    rect(s, 0, Inches(1.15), SW, Pt(4), YELLOW)
    txt(s, Inches(0.5), Inches(0.12), Inches(12), Inches(0.4),
        [[(kicker, 11, YELLOW, True)]])
    txt(s, Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.7),
        [[(title, 26, WHITE, True)]])


def pagefoot(s, n):
    txt(s, Inches(0.5), Inches(7.05), Inches(8), Inches(0.35),
        [[("PT Vale Indonesia — www.valeindonesia.com", 9, GREY, False)]])
    txt(s, Inches(12.3), Inches(7.05), Inches(0.8), Inches(0.35),
        [[(str(n), 9, GREY, False)]], align=PP_ALIGN.RIGHT)


# ═══ 1. TITLE ═════════════════════════════════════════════════════════
s = slide()
rect(s, 0, 0, SW, SH, TEAL)
rect(s, 0, Inches(4.75), SW, Pt(5), YELLOW)
txt(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.6),
    [[("Website Update Workflow", 44, WHITE, True)]])
txt(s, Inches(0.9), Inches(3.7), Inches(11.5), Inches(0.9),
    [[("How the COMMs team requests changes to ", 20, LIGHT, False),
      ("www.valeindonesia.com", 20, YELLOW, True)]])
txt(s, Inches(0.9), Inches(5.1), Inches(11.5), Inches(0.9),
    [[("Rollout briefing — COMMs & IT / GDI", 16, LIGHT, False)],
     [("September 2026", 13, LIGHT, False)]])

# ═══ 2. WHY / CONTEXT ═════════════════════════════════════════════════
s = slide(); header(s, "Why this matters")
txt(s, Inches(0.6), Inches(1.5), Inches(12), Inches(0.9),
    [[("The website is now LIVE on ", 18, DARK, False),
      ("www.valeindonesia.com", 18, TEAL, True),
      (". Keeping it current is a shared, repeatable process — not one-off fixes.", 18, DARK, False)]])
bullets(s, Inches(0.7), Inches(2.6), Inches(12), Inches(4),
    ["The site is a static mirror — fast, secure, low-cost, no CMS to log into.",
     "COMMs owns the content; GDI/IT owns the technical build & publishing.",
     "This workflow gives everyone one clear path: request → review → approve → publish.",
     "No COMMs access to servers or code is needed — you author, GDI deploys.",
     "Every change is previewed on the dev site and approved before it goes live.",
     "This is a temporary solution — PTVI should prepare for the real (long-term) solution as soon as possible."],
    size=16, gap=10)
pagefoot(s, 2)

# ═══ 3. THE MODEL (one line) ══════════════════════════════════════════
s = slide(); header(s, "The model in one line")
box_y = Inches(2.6); bw = Inches(2.2); bh = Inches(1.6); gap = Inches(0.4)
steps = [("COMMs", "authors the content", TEAL),
         ("GDI / IT", "converts & builds", TEAL),
         ("Dev preview", "review on the dev site", YELLOW),
         ("Approver", "approves the preview", TEAL),
         ("GDI / IT", "publishes to live", TEAL_D)]
x = Inches(0.36)
for i, (t, sub, col) in enumerate(steps):
    b = rect(s, x, box_y, bw, bh, col)
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = t; r.font.size = Pt(16); r.font.bold = True
    r.font.color.rgb = WHITE if col != YELLOW else DARK; r.font.name = "Calibri"
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(11)
    r2.font.color.rgb = WHITE if col != YELLOW else DARK; r2.font.name = "Calibri"
    if i < len(steps) - 1:
        txt(s, x + bw, box_y, gap, bh, [[("→", 24, GREY, True)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    x = x + bw + gap
txt(s, Inches(0.6), Inches(4.7), Inches(12), Inches(1),
    [[("COMMs never touches Azure, HTML, or code. You fill in a simple form; GDI does the rest.",
       16, GREY, False)]], align=PP_ALIGN.CENTER)
pagefoot(s, 3)

# ═══ 4. WHAT YOU CAN REQUEST ══════════════════════════════════════════
s = slide(); header(s, "What you can request")
cards = [
    ("News article", "Publish a news story (bilingual), with cover + photos.", TEAL),
    ("Document", "Add a PDF report to the Document Library.", TEAL),
    ("Edit a page", "Change text, stats, or photos on an existing page.", TEAL),
    ("New page", "Create a new page based on an existing one.", YELLOW),
]
cw = Inches(2.95); ch = Inches(2.7); cx = Inches(0.55); cy = Inches(1.8); cgap = Inches(0.2)
for (t, d, col) in cards:
    rect(s, cx, cy, cw, Inches(0.7), col)
    tt = s.shapes.add_textbox(cx, cy, cw, Inches(0.7)); tfr = tt.text_frame
    tfr.vertical_anchor = MSO_ANCHOR.MIDDLE; p = tfr.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = t; r.font.bold = True; r.font.size = Pt(16)
    r.font.color.rgb = WHITE if col != YELLOW else DARK; r.font.name = "Calibri"
    rect(s, cx, cy + Inches(0.7), cw, ch - Inches(0.7), LIGHT)
    txt(s, cx + Inches(0.18), cy + Inches(0.9), cw - Inches(0.36), ch - Inches(1),
        [[(d, 13, DARK, False)]])
    cx = cx + cw + cgap
txt(s, Inches(0.6), Inches(4.9), Inches(12), Inches(1.4),
    [[("Most updates are news & documents (quick, form-driven). ", 15, DARK, False),
      ("Page edits and new pages are done by GDI from your request.", 15, GREY, False)]])
pagefoot(s, 4)

# ═══ 5. HOW TO SUBMIT (SharePoint) ════════════════════════════════════
s = slide(); header(s, "How to submit a request")
txt(s, Inches(0.6), Inches(1.4), Inches(12), Inches(0.5),
    [[("Everything runs through one SharePoint area: ", 16, DARK, False),
      ("Website-Update-Requests", 16, TEAL, True)]])
# folder flow
folders = [("_Templates", "Copy a template from here", GREY),
           ("1-Submitted", "You drop your request folder here", TEAL),
           ("2-In-Progress", "GDI is working on it", YELLOW),
           ("3-Done-Deployed", "Live — with the links", TEAL_D)]
fy = Inches(2.2)
for (name, desc, col) in folders:
    rect(s, Inches(0.7), fy, Inches(3.2), Inches(0.7), col)
    tb = s.shapes.add_textbox(Inches(0.7), fy, Inches(3.2), Inches(0.7))
    tfr = tb.text_frame; tfr.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tfr.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = name; r.font.bold = True; r.font.size = Pt(15)
    r.font.color.rgb = WHITE if col != YELLOW else DARK; r.font.name = "Calibri"
    txt(s, Inches(4.1), fy, Inches(8), Inches(0.7), [[(desc, 14, DARK, False)]],
        anchor=MSO_ANCHOR.MIDDLE)
    fy = fy + Inches(0.85)
txt(s, Inches(0.7), Inches(5.9), Inches(12), Inches(1),
    [[("One folder per request, named  ", 13, GREY, False),
      ("YYYY-MM-DD_type_short-name", 13, TEAL, True),
      ("   (e.g. 2026-09-10_news_pkb-negotiation).", 13, GREY, False)],
     [("Put the filled template + every image/PDF you reference inside that folder.", 13, GREY, False)]])
pagefoot(s, 5)

# ═══ 6. NEW PAGE FROM EXISTING (the new capability) ═══════════════════
s = slide(); header(s, "New page — the easy way")
txt(s, Inches(0.6), Inches(1.4), Inches(12.2), Inches(0.8),
    [[("Want a new page that looks like an existing one? You edit ", 16, DARK, False),
      ("only the content", 16, TEAL, True),
      (" — never the code.", 16, DARK, False)]])
steps2 = [
    ("1", "You name the existing page you want to base it on."),
    ("2", "GDI sends you an editable page — double-click to open it in your browser."),
    ("3", "Click text to type, click images to swap, then hit “Save changes”."),
    ("4", "GDI builds the new page with the correct layout and publishes it."),
]
yy = Inches(2.5)
for (n, d) in steps2:
    c = rect(s, Inches(0.7), yy, Inches(0.6), Inches(0.6), YELLOW)
    cf = c.text_frame; cf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = cf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = n; r.font.bold = True; r.font.size = Pt(18); r.font.color.rgb = DARK
    txt(s, Inches(1.5), yy, Inches(11), Inches(0.6), [[(d, 16, DARK, False)]],
        anchor=MSO_ANCHOR.MIDDLE)
    yy = yy + Inches(0.85)
txt(s, Inches(0.7), Inches(6.1), Inches(12), Inches(0.8),
    [[("The page keeps its exact design — you can't accidentally break the layout.", 14, GREY, False)]])
pagefoot(s, 6)

# ═══ 7. ROLES, APPROVAL, SLA ══════════════════════════════════════════
s = slide(); header(s, "Roles, approval & turnaround")
# roles
txt(s, Inches(0.6), Inches(1.35), Inches(6), Inches(0.4), [[("Who does what", 16, TEAL, True)]])
bullets(s, Inches(0.7), Inches(1.85), Inches(6), Inches(2.5),
    ["COMMs — writes content, supplies images, requests changes",
     "GDI / IT — converts, builds, previews, deploys",
     "Approver — Maman Ashari Hasan Tjokke",
     ("maman.hasan@valeindonesia.com", 1),
     "Approver signs off on the dev preview before go-live"],
    size=14, gap=8)
# SLA table
txt(s, Inches(7.0), Inches(1.35), Inches(6), Inches(0.4), [[("Turnaround (once GDI starts)", 16, TEAL, True)]])
sla = [("News article", "1 business day"),
       ("Document", "1 business day"),
       ("Existing-page edit", "1–2 business days"),
       ("New page / re-clone", "Scheduled")]
ty = Inches(1.9)
for i, (a, b) in enumerate(sla):
    bg = LIGHT if i % 2 == 0 else WHITE
    rect(s, Inches(7.0), ty, Inches(5.6), Inches(0.6), bg)
    txt(s, Inches(7.15), ty, Inches(3.4), Inches(0.6), [[(a, 14, DARK, False)]], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(10.4), ty, Inches(2.1), Inches(0.6), [[(b, 14, TEAL_D, True)]], anchor=MSO_ANCHOR.MIDDLE)
    ty = ty + Inches(0.62)
txt(s, Inches(7.0), Inches(4.6), Inches(5.6), Inches(1),
    [[("Clock starts when GDI picks up the request and assumes the approver signs off promptly.",
       11, GREY, False)]])
pagefoot(s, 7)

# ═══ 8. WHAT'S ALREADY LIVE (credibility) ═════════════════════════════
s = slide(); header(s, "Already delivered")
bullets(s, Inches(0.7), Inches(1.7), Inches(12), Inches(4.5),
    ["Site live on www.valeindonesia.com (behind Azure Front Door + WAF, HTTPS).",
     "News & Document Library — fully data-driven, easy to add to.",
     "First COMMs request done end-to-end: homepage domain-change popup (Maman Ashari Hasan Tjokke).",
     "Homepage scroll button + footer alignment fixed site-wide.",
     "Request tools built: news converter + new-page content extract/inject.",
     "SharePoint request folders + templates + guides ready to go."],
    size=16, gap=12)
pagefoot(s, 8)

# ═══ 9. NEXT STEPS ════════════════════════════════════════════════════
s = slide(); header(s, "Next steps")
bullets(s, Inches(0.7), Inches(1.7), Inches(12), Inches(3.5),
    ["IT/GDI: create the SharePoint Website-Update-Requests area (package is ready to upload).",
     "COMMs: try one real request end-to-end (a news item is easiest).",
     "Confirm the approver flow with Maman Ashari Hasan Tjokke on the first go-live.",
     "Share the how-to guide with the wider COMMs team."],
    size=17, gap=14)
rect(s, Inches(0.7), Inches(5.4), Inches(11.9), Inches(1.1), LIGHT)
txt(s, Inches(1.0), Inches(5.55), Inches(11.3), Inches(0.9),
    [[("Questions or a request to start? ", 15, DARK, True),
      ("Contact GDI / IT — we'll walk you through your first one.", 15, DARK, False)]],
    anchor=MSO_ANCHOR.MIDDLE)
pagefoot(s, 9)

# ═══ 10. CLOSING ══════════════════════════════════════════════════════
s = slide()
rect(s, 0, 0, SW, SH, TEAL)
rect(s, 0, Inches(3.3), SW, Pt(4), YELLOW)
txt(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(1),
    [[("Thank you", 40, WHITE, True)]])
txt(s, Inches(0.9), Inches(3.7), Inches(11.5), Inches(1),
    [[("One simple form → your update is live. COMMs authors, GDI publishes.", 18, LIGHT, False)]])

out = "docs/COMMs-Website-Update-Workflow.pptx"
prs.save(out)
print("saved:", out, "slides:", len(prs.slides._sldIdLst))
